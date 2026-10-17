from io import BytesIO

import pytest
from openpyxl import Workbook

from watcher.source import SourceClient


def make_workbook():
    workbook = Workbook()
    workbook.active["A1"] = "Facility ID"

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    return buffer.getvalue()


class FakeResponse:
    def __init__(self, status_code, content=b"", url="https://example.com/report"):
        self.status_code = status_code
        self.content = content
        self.url = url
        self.text = content.decode("utf-8", errors="replace")

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            raise requests.HTTPError(f"HTTP {self.status_code}")


class FakeSession:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return next(self.responses)


def test_rejects_corrupt_workbook():
    with pytest.raises(ValueError, match="valid workbook"):
        SourceClient._validate_workbook(b"not an Excel file")


def test_accepts_valid_workbook():
    SourceClient._validate_workbook(make_workbook())


def test_download_caches_valid_workbook(tmp_path, monkeypatch):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    response = FakeResponse(200, make_workbook())
    session = FakeSession([response])
    client.session = session

    url = "https://example.com/report"

    first_path = client.download_workbook(url)
    second_path = client.download_workbook(url)

    assert first_path == second_path
    assert first_path.exists()
    assert first_path.read_bytes() == response.content
    assert len(session.calls) == 1


def test_retries_temporary_http_error(tmp_path, monkeypatch):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    session = FakeSession([
        FakeResponse(503),
        FakeResponse(200, make_workbook()),
    ])
    client.session = session
    monkeypatch.setattr("watcher.source.time.sleep", lambda _: None)

    response = client._request("https://example.com/report")

    assert response.status_code == 200
    assert len(session.calls) == 2


def test_403_retries_with_alternate_user_agent(tmp_path, monkeypatch):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    session = FakeSession([
        FakeResponse(403),
        FakeResponse(200, b"page"),
    ])
    client.session = session
    monkeypatch.setattr("watcher.source.time.sleep", lambda _: None)

    response = client._request("https://example.com")

    assert response.status_code == 200
    assert len(session.calls) == 2

    first_agent = session.calls[0][1]["headers"]["User-Agent"]
    second_agent = session.calls[1][1]["headers"]["User-Agent"]

    assert first_agent != second_agent


def test_rejects_cross_host_redirect(tmp_path, monkeypatch):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    client.session = FakeSession([
        FakeResponse(
            200,
            b"redirected page",
            url="https://unrelated.example.org/page",
        )
    ])

    with pytest.raises(ValueError, match="redirected to another website"):
        client.fetch_page("https://example.com")

def test_rejects_workbook_that_fails_during_reading(monkeypatch):
    class BrokenSheet:
        title = "Sheet1"

        def iter_rows(self, values_only=False):
            yield ("Facility ID",)
            raise OSError("Simulated truncated worksheet")

    class BrokenWorkbook:
        sheetnames = ["Sheet1"]
        worksheets = [BrokenSheet()]

        def close(self):
            pass

    def fake_load_workbook(*args, **kwargs):
        return BrokenWorkbook()

    monkeypatch.setattr(
        "watcher.source.load_workbook",
        fake_load_workbook,
    )

    with pytest.raises(ValueError, match="valid.*workbook"):
        SourceClient._validate_workbook(b"fake workbook bytes")

def test_rejects_cross_host_workbook_redirect(tmp_path):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    client.session = FakeSession([
        FakeResponse(
            200,
            make_workbook(),
            url="https://unrelated.example.org/report",
        )
    ])

    with pytest.raises(
        ValueError,
        match="Workbook URL redirected to another website",
    ):
        client.download_workbook("https://example.com/report")

    assert list(tmp_path.iterdir()) == []


def test_404_is_reported_without_retries(tmp_path, monkeypatch):
    client = SourceClient(cache_dir=tmp_path, min_interval=0)
    session = FakeSession([FakeResponse(404)])
    client.session = session

    monkeypatch.setattr("watcher.source.time.sleep", lambda _: None)

    with pytest.raises(RuntimeError, match="HTTP request failed: 404"):
        client._request("https://example.com/missing")

    assert len(session.calls) == 1