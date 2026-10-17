"""Workbook parsing with dynamic header detection and project extraction."""

from pathlib import Path
from openpyxl import load_workbook


def read_workbook(path: str | Path, config: dict, allow_missing: bool = False) -> dict:
    """Read project records from the configured worksheet using dynamic header detection."""
    workbook_path = Path(path)

    # load_workbook in read_only mode is memory efficient for large files (e.g. 20MB+)
    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet_name = config["sheet"]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {sheet_name!r} not found. "
                f"Available worksheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]
        required_columns = config["required_columns"]
        id_column = config["id_column"]

        # Search for header candidates among the first 60 rows
        candidates = []
        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            if row_number > 60:
                break

            normalized = [
                str(value).strip() if value is not None else ""
                for value in row
            ]

            if id_column in normalized:
                exact_req_matches = sum(
                    1 for col in required_columns if col in normalized
                )
                total_named_cols = sum(1 for val in normalized if val)
                candidates.append((exact_req_matches, total_named_cols, row_number, normalized))

        if not candidates:
            raise ValueError(
                f"Could not find a header row containing ID column {id_column!r}."
            )

        # Pick candidate with most required column matches, breaking ties by total named columns
        candidates.sort(key=lambda c: (c[0], c[1]), reverse=True)
        best_req_matches, _, header_row, headers = candidates[0]

        if best_req_matches == 0 and not allow_missing:
            raise ValueError(
                "Could not find a header row matching project ID and required columns."
            )

        # Map each available non-empty column name to its column index
        column_indexes = {
            name: index
            for index, name in enumerate(headers)
            if name
        }

        if id_column not in column_indexes:
            raise ValueError(f"ID column {id_column!r} not indexed in header row.")

        id_index = column_indexes[id_column]

        columns_to_read = (
            list(column_indexes.keys())
            if allow_missing
            else [id_column, *[c for c in required_columns if c in column_indexes]]
        )

        records = []
        # Filter headers or subheaders like 'ID' in subheader rows
        skip_ids = {id_column.casefold(), "id", "none", "null", "facility id"}

        for row in sheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            if id_index >= len(row):
                continue

            raw_id = row[id_index]
            if raw_id is None:
                continue

            project_id = str(raw_id).strip()
            if not project_id or project_id.casefold() in skip_ids:
                continue

            record = {"project_id": project_id}
            for col in columns_to_read:
                if col == "project_id":
                    continue
                col_i = column_indexes[col]
                record[col] = row[col_i] if col_i < len(row) else None

            records.append(record)

        return {
            "sheet": sheet_name,
            "header_row": header_row,
            "columns": headers,
            "projects": len(records),
            "records": records,
        }

    finally:
        workbook.close()
